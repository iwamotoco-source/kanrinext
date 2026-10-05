"""Independent content checks on outputs from test-compression.cjs."""
import argparse
import json
import random
import subprocess
import zipfile
from pathlib import Path
from PIL import Image, ImageChops, ImageStat
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument('--fixtures', default='/tmp/compression-test')
parser.add_argument('--prepare', action='store_true')
args = parser.parse_args()
p = Path(args.fixtures)
p.mkdir(parents=True, exist_ok=True)
if args.prepare:
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill
    rng = random.Random(7)
    im = Image.new('RGB', (640, 480))
    im.putdata([(int(x/640*180)+rng.randrange(40), int(y/480*180)+rng.randrange(40), 60+rng.randrange(80)) for y in range(480) for x in range(640)])
    im.save(p/'photo.jpg', quality=100, subsampling=0)
    Image.new('RGBA', (320, 240), (20, 90, 100, 120)).save(p/'transparent.png', compress_level=0)
    w = openpyxl.Workbook()
    s = w.active
    s.title = '測定表'
    s['A1'] = '測定結果'
    s['A2'], s['A3'], s['A4'] = 10, 20, '=SUM(A2:A3)'
    s['A1'].font = Font(bold=True, color='FF112233')
    s['A1'].fill = PatternFill('solid', fgColor='FFABCDEF')
    s.add_image(XLImage(p/'photo.jpg'), 'C2')
    w.save(p/'fixture.xlsx')
    subprocess.run(['ffmpeg', '-hide_banner', '-loglevel', 'error', '-y', '-f', 'lavfi', '-i', 'testsrc2=size=320x240:rate=15', '-f', 'lavfi', '-i', 'sine=frequency=440', '-t', '2', '-c:v', 'libx264', '-crf', '0', '-c:a', 'aac', '-b:a', '192k', str(p/'video.mp4')], check=True)
    print('Fixtures prepared')
    raise SystemExit

a, b = Image.open(p/'transparent.png'), Image.open(p/'transparent_compressed.png')
assert a.size == b.size and a.tobytes() == b.tobytes(), 'PNG pixels/alpha changed'
a, b = Image.open(p/'photo.jpg'), Image.open(p/'photo_compressed.jpg')
assert a.size == b.size
delta = sum(ImageStat.Stat(ImageChops.difference(a, b)).mean)/3
# Bound mean channel error to 5% of the 8-bit range on this noisy fixture.
assert delta < 255*.05, f'JPEG mean absolute channel error too large: {delta}'
w = openpyxl.load_workbook(p/'fixture_compressed.xlsx')
s = w['測定表']
assert s['A4'].value == '=SUM(A2:A3)' and s['A1'].font.bold
assert s['A1'].fill.fgColor.rgb == 'FFABCDEF' and len(s._images) == 1
with zipfile.ZipFile(p/'fixture.xlsx') as a, zipfile.ZipFile(p/'fixture_compressed.xlsx') as b:
    assert a.namelist() == b.namelist()
    assert all(a.read(n) == b.read(n) for n in a.namelist() if not n.startswith('xl/media/'))
with zipfile.ZipFile(p/'macro_compressed.xlsm') as z:
    assert z.read('xl/vbaProject.bin') == bytes([1, 2, 3, 4, 5])
with zipfile.ZipFile(p/'batch.zip') as z:
    assert len(z.namelist()) == 2 and z.testzip() is None

import fitz
a, b = fitz.open(p/'source.pdf'), fitz.open(p/'drawing_compressed.pdf')
assert len(a) == len(b) == 1 and a[0].rect == b[0].rect
assert a[0].get_text() == b[0].get_text() and 'Voltage 200 V' in b[0].get_text()
assert [w.field_value for w in b[0].widgets()] == ['ABC123']
streams = []
for name in ['video.mp4', 'video_compressed.mp4']:
    streams.append(json.loads(subprocess.check_output(['ffprobe', '-v', 'error', '-show_streams', '-of', 'json', str(p/name)]))['streams'])
a, b = streams
for key in ['width', 'height', 'r_frame_rate', 'nb_frames']:
    assert a[0][key] == b[0][key], key
assert len(a) == len(b) == 2
hashes = [subprocess.check_output(['ffmpeg', '-hide_banner', '-loglevel', 'error', '-i', str(p/name), '-map', '0:a', '-c', 'copy', '-f', 'hash', '-hash', 'sha256', '-']).strip() for name in ['video.mp4', 'video_compressed.mp4']]
assert hashes[0] == hashes[1], 'Encoded audio changed'
print(f'PASS: PNG pixels/alpha; JPEG mean error {delta:.2f}/255; Excel formulas/style/XML/VBA; PDF text/form/pages; video dimensions/fps/frames; exact audio; ZIP CRC')
